"""Integration tests for PDF to CSV / Excel conversion."""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pymupdf as fitz
import pytest

from pdf_tool.adapters.pdf_to_csv import PdfToCsvAdapter
from pdf_tool.exceptions import ConversionFailedError


def _create_pdf_with_table(pdf_path: Path) -> None:
    """Helper to generate a PDF containing a structured table."""
    doc = fitz.open()
    p = doc.new_page()

    # Draw table bounding boxes and cell divider lines
    p.draw_rect(fitz.Rect(50, 50, 350, 150), color=(0, 0, 0), width=1)
    p.draw_line(fitz.Point(50, 80), fitz.Point(350, 80), color=(0, 0, 0), width=1)
    p.draw_line(fitz.Point(200, 50), fitz.Point(200, 150), color=(0, 0, 0), width=1)

    p.insert_text((60, 70), "Item", fontsize=11)
    p.insert_text((210, 70), "Cost", fontsize=11)
    p.insert_text((60, 110), "Coffee", fontsize=11)
    p.insert_text((210, 110), "5", fontsize=11)

    doc.save(str(pdf_path))
    doc.close()


def test_pdf_to_csv(tmp_path: Path) -> None:
    """PDF table extracts accurately into a CSV file."""
    pdf_path = tmp_path / "table.pdf"
    _create_pdf_with_table(pdf_path)

    csv_path = tmp_path / "table.csv"
    adapter = PdfToCsvAdapter()
    result = adapter.convert(pdf_path, csv_path)

    assert result.output_path.exists()
    assert result.engine == "pdfplumber"

    with csv_path.open(newline="", encoding="utf-8") as f:
        reader = list(csv.reader(f))
        assert any("Item" in row for row in reader)
        assert any("Coffee" in row for row in reader)


def test_pdf_to_xlsx(tmp_path: Path) -> None:
    """PDF table extracts accurately into an Excel workbook."""
    pdf_path = tmp_path / "table.pdf"
    _create_pdf_with_table(pdf_path)

    xlsx_path = tmp_path / "table.xlsx"
    adapter = PdfToCsvAdapter()
    result = adapter.convert(pdf_path, xlsx_path)

    assert result.output_path.exists()
    assert result.engine == "pdfplumber"

    wb = openpyxl.load_workbook(str(xlsx_path))
    assert len(wb.sheetnames) >= 1
    sheet = wb[wb.sheetnames[0]]
    cell_values = [cell.value for row in sheet.rows for cell in row if cell.value]
    assert "Item" in cell_values
    assert "Coffee" in cell_values


def test_pdf_without_tables_raises(tmp_path: Path) -> None:
    """PDF with empty pages raises ConversionFailedError."""
    empty_pdf = tmp_path / "empty.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(empty_pdf))
    doc.close()

    csv_path = tmp_path / "out.csv"
    adapter = PdfToCsvAdapter()
    with pytest.raises(ConversionFailedError) as exc_info:
        adapter.convert(empty_pdf, csv_path)
    assert "No tabular data" in str(exc_info.value)
